# -*- coding: utf-8 -*-
"""اختبارات الاستيراد من Excel ووضع الشبكة (تعمل على الخادم الحقيقي)."""
import os, sys, io, datetime, sqlite3, unittest, urllib.request, urllib.error, threading, shutil, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from test_acceptance import Base, Client, PW
import server

D = datetime.datetime
HDR1 = ['اسم المراجع ', 'تاريخ اخر المراجعة ', 'نوع العمل ', 'رقم السن', 'المدفوع ', 'المبلغ المتبقي', 'الملاحظات', 'عمود8']
HDR2 = ['اسم المراجع ', 'تاريخ اخر المراجعة ', 'نوع العمل ', 'رقم السن', 'المدفوع ', 'الباقي', 'الملاحظات', None]


def make_xlsx(sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
        ws.cell(row=1, column=100).value = 'x'      # عمود بعيد يجب تجاهله
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


ROWS_2022 = [HDR1, HDR2,
    ['علي حسن ', '25/1', 'ج1', 'LL6', 100, 400, 'mb=20', None],              # تكلفة 500، متبقي 400
    ['علي  حسن', '26/1', 'تسديد', None, 150, None, None, None],               # دفعة فقط → متبقي 250
    ['عليّ حسن', '27/1', 'ضوئية', 'ur123', 100, 150, None, None],              # تكلفة جديدة 0 + متبقي 150 (الاسم بتشكيل)
    ['علي حسن', '27/1', 'ضوئية', 'ur123', 100, 150, None, None],              # مكرر تماماً
    ['زينب علي', '13/7', 'قلع', 'lre', 20, None, None, None],
    ['زينب علي', D(2022, 3, 7), 'ج4', 'ul3+4', 30, None, None, None],          # حقيقته 3 تموز (مقلوب)
    ['زينب علي', '15/7', 'تنظيف', None, 15, None, None, None],
    ['زينب علي', D(2022, 7, 5), 'تنظيف', 'u11', 15, None, None, None],         # 5 تموز فعلاً
    ['سالب', '5/2', 'ج1', 'ul6', -250, None, None, None],                     # مرفوض
    [None, '5/2', 'ج1', 'ul6', 25, None, None, None],                         # بدون اسم
    ['بدون تاريخ', None, 'ج1', 'ul6', 25, None, None, None],
    ['اسماء احمد', '1/2', 'فحص', 'zz9', 10, None, 'ملاحظة', 20.5],            # اسم يبدأ بـ «اسم» ليس عنواناً
]
ROWS_2023 = [['اسم المراجع', 'التاريخ', 'نوع العمل', 'رقم السن', 'المدفوع', 'المتبقي', 'عمود7', 'عمود8'],
    ['علي حسن', D(2023, 1, 2), 'ج2', 'lr6', 50, None, 'مرجع', None]]


class T20_Import(Base):
    def preview(self, data=None, mult=1000, name='x.xlsx'):
        data = data or make_xlsx({'2022': ROWS_2022, '2023': ROWS_2023})
        return self.c.ok('POST', '/api/import/preview?mult=%d&filename=%s' % (mult, name), raw=data)

    def test_preview_and_commit(self):
        pre = self.preview()
        s = pre['summary']
        self.assertEqual((s['duplicates_in_file'], s['rows_skipped'], s['patients']), (1, 3, 3))
        self.assertEqual(s['skipped_by_reason'], {'بدون اسم صالح': 1, 'مبلغ مدفوع سالب': 1, 'بدون تاريخ': 1})
        self.assertGreaterEqual(s['date_fixed'], 1)
        self.assertEqual(s['total_paid'], (100 + 150 + 100 + 20 + 30 + 15 + 15 + 10 + 50) * 1000)
        self.assertEqual(self.c.call('GET', '/api/patients')[1]['patients'], [])     # المعاينة لا تُدخل شيئاً
        r = self.c.ok('POST', '/api/import/commit', {'token': pre['token']})
        self.assertTrue(r['backup'].startswith('pre_import_'))
        res = self.c.ok('GET', '/api/patients?q=' + urllib.parse.quote('علي حسن'))['patients']
        self.assertEqual(len(res), 1)                              # كل صيغ الاسم دُمجت بملف واحد
        pid = res[0]['id']
        full = self.c.ok('GET', '/api/patients/%d' % pid)
        # الرصيد: 500 تكلفة − 100 −150 دفعات ... = آخر «متبقي» مكتوب (150) ثم دفعة 50 لاحقة بلا متبقٍ = 150
        self.assertEqual(full['summary']['remaining'], 150000)
        self.assertEqual(full['summary']['paid'], (100 + 150 + 100 + 50) * 1000)
        teeth = {p['tooth_no'] for v in full['visits'] for p in v['procedures']}
        self.assertIn('46', teeth); self.assertIn('36', teeth)      # LL6 → 36 ، lr6 → 46
        multi = [p for v in full['visits'] for p in v['procedures'] if p['name'] == 'ضوئية'][0]
        self.assertIsNone(multi['tooth_no']); self.assertIn('11، 12، 13', multi['details'])
        self.assertTrue(any(p['name'] == 'ج1' and 'mb=20' in (p['details'] or '') for v in full['visits'] for p in v['procedures']))
        self.assertEqual(len(full['payments']), 4)                  # التسديد دفعة بدون زيارة
        self.assertEqual(sorted(v['visit_date'] for v in full['visits'])[0], '2022-01-25')
        # زينب: تاريخ مقلوب صُحّح إلى 3 تموز، وسن الأطفال lre → 85
        z = self.c.ok('GET', '/api/patients?q=' + urllib.parse.quote('زينب'))['patients'][0]
        zf = self.c.ok('GET', '/api/patients/%d' % z['id'])
        self.assertIn('2022-07-03', [v['visit_date'] for v in zf['visits']])
        self.assertNotIn('2022-03-07', [v['visit_date'] for v in zf['visits']])
        self.assertIn('2022-07-05', [v['visit_date'] for v in zf['visits']])
        self.assertIn('85', {p['tooth_no'] for v in zf['visits'] for p in v['procedures']})
        self.assertEqual(zf['summary']['remaining'], 0)
        # اسم يبدأ بـ «اسم» استُورد كمراجع، وسن غير مفهوم حُفظ نصاً
        a = self.c.ok('GET', '/api/patients?q=' + urllib.parse.quote('اسماء'))['patients'][0]
        af = self.c.ok('GET', '/api/patients/%d' % a['id'])
        self.assertIn('zz9', af['visits'][0]['procedures'][0]['details']); self.assertIn('20.5', af['visits'][0]['procedures'][0]['details'])
        # تدقيق ودفاتر
        rep = self.c.ok('GET', '/api/reports/period?from=2022-01-01&to=2023-12-31')['stats']
        self.assertEqual(rep['receipts'], s['total_paid'])
        self.assertEqual(self.c.ok('GET', '/api/audit?entity=import')['rows'][0]['action'], 'import')

    def test_reimport_is_idempotent_and_merges_with_existing(self):
        existing = self.patient('علي حسن', '07700001111')
        pre = self.preview()
        self.assertEqual(pre['summary']['patients_existing'], 1)
        self.c.ok('POST', '/api/import/commit', {'token': pre['token']})
        self.assertEqual(len(self.c.ok('GET', '/api/patients?q=' + urllib.parse.quote('علي حسن'))['patients']), 1)
        self.assertEqual(len(self.c.ok('GET', '/api/patients/%d' % existing['id'])['payments']), 4)
        again = self.preview()
        self.assertEqual((again['summary']['rows_imported'], again['summary']['duplicates_previous']), (0, 10))
        s, d = self.c.call('POST', '/api/import/commit', {'token': again['token']})
        self.assertEqual(s, 400)
        n = len(self.c.ok('GET', '/api/patients?limit=500')['patients'])
        self.assertEqual(n, 3)
        # ملف إضافي فيه صف جديد وصف قديم → يدخل الجديد فقط
        more = make_xlsx({'2022': ROWS_2022 + [['علي حسن', '30/1', 'ج3', 'ur6', 10, None, None, None]], '2023': ROWS_2023})
        p3 = self.preview(more)
        self.assertEqual(p3['summary']['rows_imported'], 1)
        self.c.ok('POST', '/api/import/commit', {'token': p3['token']})
        self.assertEqual(len(self.c.ok('GET', '/api/patients/%d' % existing['id'])['payments']), 5)

    def test_multiplier_bad_file_and_permissions(self):
        pre = self.preview(mult=1)
        self.assertEqual(pre['summary']['total_paid'], 100 + 150 + 100 + 20 + 30 + 15 + 15 + 10 + 50)
        s, d = self.c.call('POST', '/api/import/preview?mult=1000', raw=b'this is not a spreadsheet at all')
        self.assertEqual(s, 200)       # يُقرأ كـ CSV ولا يجد عناوين، فلا صفوف صالحة
        self.assertEqual(d['summary']['rows_imported'], 0)
        self.assertEqual(self.c.call('POST', '/api/import/preview?mult=7', raw=b'x')[0], 400)
        self.c.ok('POST', '/api/users', {'username': 'rec9', 'password': PW, 'role': 'reception'})
        rec = Client(self.srv.server_address[1]); rec.login('rec9')
        self.assertEqual(rec.call('POST', '/api/import/preview?mult=1000', raw=make_xlsx({'2022': ROWS_2022}))[0], 403)
        self.assertEqual(rec.call('POST', '/api/import/commit', {'token': 'x'})[0], 403)

    def test_csv_import(self):
        csv_bytes = ('اسم المراجع,التاريخ,نوع العمل,رقم السن,المدفوع,المتبقي\nسارة كريم,2024-03-05,حشوة,ur6,25,\nسارة كريم,2024-03-05,حشوة,ur6,25,\n').encode('utf-8-sig')
        pre = self.c.ok('POST', '/api/import/preview?mult=1000&filename=a.csv', raw=csv_bytes)
        self.assertEqual((pre['summary']['rows_imported'], pre['summary']['duplicates_in_file']), (1, 1))
        self.c.ok('POST', '/api/import/commit', {'token': pre['token']})
        self.assertEqual(self.c.ok('GET', '/api/reports/debtors')['total_dues'], 0)


class T21_LanMode(unittest.TestCase):
    def test_lan_mode_host_rules(self):
        tmp = tempfile.mkdtemp(prefix='clinic_lan_')
        srv = server.make_server(tmp, 0, lan=True)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        port = srv.server_address[1]
        try:
            c = Client(port); c.setup()
            s = c.ok('GET', '/api/session')
            self.assertTrue(s['lan'])
            for host, code in [('192.168.1.20:%d' % port, 200), ('10.0.0.5:%d' % port, 200), ('localhost:%d' % port, 200), ('evil.example.com', 403), ('8.8.8.8', 403)]:
                req = urllib.request.Request('http://127.0.0.1:%d/api/session' % port, headers={'Host': host})
                try:
                    code_got = urllib.request.urlopen(req).status
                except urllib.error.HTTPError as e:
                    code_got = e.code
                self.assertEqual(code_got, code, host)
            # بدون وضع الشبكة: عنوان خاص مرفوض
            tmp2 = tempfile.mkdtemp(prefix='clinic_nolan_')
            srv2 = server.make_server(tmp2, 0)
            threading.Thread(target=srv2.serve_forever, daemon=True).start()
            req = urllib.request.Request('http://127.0.0.1:%d/api/session' % srv2.server_address[1], headers={'Host': '192.168.1.20'})
            with self.assertRaises(urllib.error.HTTPError) as cm:
                urllib.request.urlopen(req)
            self.assertEqual(cm.exception.code, 403)
            self.assertEqual(srv2.server_address[0], '127.0.0.1'); self.assertEqual(srv.server_address[0], '0.0.0.0')
            srv2.shutdown(); srv2.server_close(); shutil.rmtree(tmp2, ignore_errors=True)
        finally:
            srv.shutdown(); srv.server_close(); shutil.rmtree(tmp, ignore_errors=True)

    def test_two_laptops_share_data_concurrently(self):
        tmp = tempfile.mkdtemp(prefix='clinic_two_')
        srv = server.make_server(tmp, 0, lan=True)
        threading.Thread(target=srv.serve_forever, daemon=True).start()
        port = srv.server_address[1]
        try:
            a, b = Client(port), Client(port); a.setup()
            a.ok('POST', '/api/users', {'username': 'lap2', 'password': PW, 'role': 'reception'})
            self.assertEqual(b.login('lap2')[0], 200)
            ids = []
            def add(cl, n):
                for i in range(15):
                    s, d = cl.call('POST', '/api/patients', {'full_name': 'مراجع %s %d' % (n, i), 'phone': '0780%07d' % (i + (1000 if n == 'ب' else 0))})
                    ids.append((s, d['patient']['file_no'] if s == 200 else d))
            ts = [threading.Thread(target=add, args=(a, 'أ')), threading.Thread(target=add, args=(b, 'ب'))]
            [t.start() for t in ts]; [t.join() for t in ts]
            self.assertTrue(all(s == 200 for s, _ in ids))
            nums = [n for _, n in ids]
            self.assertEqual(len(set(nums)), 30)                       # أرقام ملفات فريدة رغم التزامن من جهازين
            self.assertEqual(len(a.ok('GET', '/api/patients?limit=500')['patients']), 30)
            self.assertEqual(len(b.ok('GET', '/api/patients?limit=500')['patients']), 30)   # الجهازان يريان نفس البيانات
        finally:
            srv.shutdown(); srv.server_close(); shutil.rmtree(tmp, ignore_errors=True)


if __name__ == '__main__':
    unittest.main(verbosity=2)
